"""Generate the PROVISIONAL P9-C7 workbook (cells, panes, comments, pictures).

    python3 -I scripts/testdata-src/xlsx/cells_c7.py OUT_DIR

Writes provisional-cells-c7.xlsx (see testdata/README.md):

* sheet `Cells`: text overflowing into empty cells and stopped by a value,
  a date in a column too narrow for it (####), wrap text, 45° rotated text,
  a right indent, horizontal alignment Fill; panes frozen while scrolled
  (`sheetView/@topLeftCell` A20, `pane ySplit=5 topLeftCell=A25`) and a
  manual row break before row 21;
* sheet `Split`: split panes (`pane xSplit=2400 ySplit=1800`, no state)
  with the top-left pane scrolled to B2 and the bottom-right one to D10;
* sheet `Pictures`: a picture "don't move or size with cells"
  (`editAs="absolute"`) cropped (`a:srcRect`), 40 % transparent
  (`a:alphaModFix`) with a dark red border (`a:ln`), and one "move but
  don't size" (`editAs="oneCell"`) above it in the z-order;
* a threaded comment on `Cells!B2` with a reply @mentioning Ann Lee
  (`mentions`), the persons with `userId` / `providerId` (AD).

openpyxl writes the workbook; the parts it cannot write (threaded comments,
persons, the picture formats, the pane top-left cells) are patched in, as
Excel writes them. openpyxl is installed per user (docs/setup.md); `-I`
drops the user site directory, so it is added back explicitly.
"""

import io
import site
import struct
import sys
import zipfile
import zlib

sys.path.append(site.getusersitepackages())

from openpyxl import Workbook  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.styles import Alignment  # noqa: E402
from openpyxl.worksheet.pagebreak import Break  # noqa: E402

ANN = "{7C1D2E3F-4A5B-4C6D-8E7F-000000000001}"
BOB = "{7C1D2E3F-4A5B-4C6D-8E7F-000000000002}"
T1 = "{7C1D2E3F-4A5B-4C6D-8E7F-0000000000A1}"
T1R = "{7C1D2E3F-4A5B-4C6D-8E7F-0000000000A2}"
M1 = "{7C1D2E3F-4A5B-4C6D-8E7F-0000000000B1}"

PLACEHOLDER = (
    "[Threaded comment]\n\nYour version of Excel allows you to read this threaded comment; "
    "however, any edits to it will get removed if the file is opened in a newer version of Excel. "
    "Learn more: https://go.microsoft.com/fwlink/?linkid=870924\n\nComment:\n    "
)


def png(w, h, rgb):
    """A plain PNG of `w × h` pixels in one colour."""
    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))
    chunk = lambda t, d: struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)  # noqa: E731
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


def fixed(wb):
    """Stable document properties (openpyxl stamps the current time)."""
    import datetime

    t = datetime.datetime(2026, 1, 1)
    wb.properties.created = t
    wb.properties.modified = t
    wb.properties.creator = "Rusuite"
    wb.properties.lastModifiedBy = "Rusuite"


def build():
    wb = Workbook()
    fixed(wb)
    ws = wb.active
    ws.title = "Cells"
    ws["A1"] = "A long text that overflows into the empty cells beside it"
    ws["D1"] = "stop"
    ws["A2"] = "Review"
    ws["B3"] = "=DATE(2024,12,31)"
    ws["B3"].number_format = "m/d/yyyy"
    ws.column_dimensions["B"].width = 4
    ws["A4"] = "Wrapped text in a cell of its own"
    ws["A4"].alignment = Alignment(wrap_text=True)
    ws["A5"] = "Turned forty-five degrees"
    ws["A5"].alignment = Alignment(text_rotation=45)
    ws["C6"] = "indented"
    ws["C6"].alignment = Alignment(horizontal="right", indent=2)
    ws["A7"] = "-="
    ws["A7"].alignment = Alignment(horizontal="fill")
    for r in range(8, 61):
        ws.cell(row=r, column=1, value=f"line {r}")
    ws.row_breaks.append(Break(id=20))
    # Excel's placeholder for the threaded comment on B2.
    ws["B2"].comment = Comment(PLACEHOLDER + "Is the date right?\nReply:\n    @Ann Lee yes, checked.", "tc=" + T1)

    split = wb.create_sheet("Split")
    for r in range(1, 41):
        for c in range(1, 9):
            split.cell(row=r, column=c, value=r * 10 + c)

    pics = wb.create_sheet("Pictures")
    pics["A1"] = "Logo below"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


THREADED = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<ThreadedComments xmlns="http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments" xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><threadedComment ref="B2" dT="2026-10-18T09:00:00.00" personId="{BOB}" id="{T1}"><text>Is the date right?</text></threadedComment><threadedComment ref="B2" dT="2026-10-18T09:30:00.00" personId="{ANN}" id="{T1R}" parentId="{T1}"><text>@Ann Lee yes, checked.</text><mentions><mention mentionpersonId="{ANN}" mentionId="{M1}" startIndex="0" length="8"/></mentions></threadedComment></ThreadedComments>"""

PERSONS = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<personList xmlns="http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments" xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><person displayName="Ann Lee" id="{ANN}" userId="ann.lee@example.com" providerId="AD"/><person displayName="Bob Stone" id="{BOB}" userId="bob@example.com" providerId="AD"/></personList>"""


def patch_sheet(body, name):
    """The pane top-left cells openpyxl cannot write."""
    if name == "xl/worksheets/sheet1.xml":
        start = body.index(b"<sheetViews>") + len(b"<sheetViews>")
        end = body.index(b"</sheetViews>")
        body = (
            body[:start]
            + b'<sheetView tabSelected="1" topLeftCell="A20" workbookViewId="0"><pane ySplit="5" topLeftCell="A25" activePane="bottomLeft" state="frozen"/><selection pane="bottomLeft" activeCell="A25" sqref="A25"/></sheetView>'
            + body[end:]
        )
    elif name == "xl/worksheets/sheet2.xml":
        start = body.index(b"<sheetViews>") + len(b"<sheetViews>")
        end = body.index(b"</sheetViews>")
        body = (
            body[:start]
            + b'<sheetView topLeftCell="B2" workbookViewId="0"><pane xSplit="2400" ySplit="1800" topLeftCell="D10" activePane="bottomRight"/><selection pane="bottomRight" activeCell="D10" sqref="D10"/></sheetView>'
            + body[end:]
        )
    return body


def anchor(edit_as, c0, r0, c1, r1, pid, name, rid, extra_blip="", extra_fill="", extra_sp=""):
    """A `xdr:twoCellAnchor` with a picture, as Excel writes it."""
    return (
        f'<xdr:twoCellAnchor editAs="{edit_as}"><xdr:from><xdr:col>{c0}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{r0}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
        f'<xdr:to><xdr:col>{c1}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{r1}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
        f'<xdr:pic><xdr:nvPicPr><xdr:cNvPr id="{pid}" name="{name}" descr="{name} logo"/><xdr:cNvPicPr><a:picLocks noChangeAspect="1"/></xdr:cNvPicPr></xdr:nvPicPr>'
        f'<xdr:blipFill><a:blip xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:embed="{rid}">{extra_blip}</a:blip>{extra_fill}<a:stretch><a:fillRect/></a:stretch></xdr:blipFill>'
        f'<xdr:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>{extra_sp}</xdr:spPr></xdr:pic><xdr:clientData/></xdr:twoCellAnchor>'
    )


DRAWING = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    + anchor("absolute", 1, 2, 3, 7, 2, "Picture 1", "rId1", '<a:alphaModFix amt="60000"/>', '<a:srcRect l="10000" r="5000" b="2500"/>',
             '<a:ln w="19050"><a:solidFill><a:srgbClr val="C00000"/></a:solidFill></a:ln>')
    + anchor("oneCell", 2, 4, 4, 9, 3, "Picture 2", "rId2")
    + "</xdr:wsDr>"
)

DRAWING_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image1.png"/>'
    '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image2.png"/>'
    "</Relationships>"
)

SHEET3_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rIdDr1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing1.xml"/>'
    "</Relationships>"
)


def finish(data):
    """Adds the threaded comments and persons, and patches panes and pictures."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "[Content_Types].xml":
                body = body.replace(
                    b"</Types>",
                    b'<Override PartName="/xl/threadedComments/threadedComment1.xml" '
                    b'ContentType="application/vnd.ms-excel.threadedcomments+xml"/>'
                    b'<Override PartName="/xl/persons/person.xml" '
                    b'ContentType="application/vnd.ms-excel.person+xml"/></Types>',
                )
            elif info.filename == "xl/worksheets/_rels/sheet1.xml.rels":
                body = body.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdTc1" Type="http://schemas.microsoft.com/office/2017/10/relationships/threadedComment" '
                    b'Target="../threadedComments/threadedComment1.xml"/></Relationships>',
                )
            elif info.filename == "xl/_rels/workbook.xml.rels":
                body = body.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdP1" Type="http://schemas.microsoft.com/office/2017/10/relationships/person" '
                    b'Target="persons/person.xml"/></Relationships>',
                )
            elif info.filename.startswith("xl/worksheets/sheet") and info.filename.endswith(".xml"):
                body = patch_sheet(body, info.filename)
            if info.filename == "[Content_Types].xml":
                body = body.replace(
                    b"</Types>",
                    b'<Default Extension="png" ContentType="image/png"/>'
                    b'<Override PartName="/xl/drawings/drawing1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>',
                )
            if info.filename == "xl/worksheets/sheet3.xml":
                body = body.replace(b"</worksheet>", b'<drawing xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rIdDr1"/></worksheet>')
            z.writestr(info, body)
        for name, text in [
            ("xl/threadedComments/threadedComment1.xml", THREADED.encode()),
            ("xl/persons/person.xml", PERSONS.encode()),
            ("xl/drawings/drawing1.xml", DRAWING.encode()),
            ("xl/drawings/_rels/drawing1.xml.rels", DRAWING_RELS.encode()),
            ("xl/worksheets/_rels/sheet3.xml.rels", SHEET3_RELS.encode()),
            ("xl/media/image1.png", png(40, 20, (0x44, 0x72, 0xC4))),
            ("xl/media/image2.png", png(20, 20, (0xED, 0x7D, 0x31))),
        ]:
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, text)
    return out.getvalue()


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else "testdata/xlsx"
    with open(f"{out}/provisional-cells-c7.xlsx", "wb") as f:
        f.write(finish(build()))


if __name__ == "__main__":
    main()
