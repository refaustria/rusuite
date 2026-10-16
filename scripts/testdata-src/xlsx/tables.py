"""Generate the PROVISIONAL tables / outline reference workbook.

    python3 -I scripts/testdata-src/xlsx/tables.py OUT_DIR

Writes provisional-tables.xlsx (see testdata/README.md): sheet `Sales` with
two Excel tables (`Sales`: header, calculated column, totals row with a
label and SUBTOTAL totals, Medium 9 with banded rows and an AutoFilter;
`Small`: Light 1, no totals) and formulas with structured references;
sheet `Outline` with nested row groups (level 2 collapsed), a column
group, summary rows above and manual page breaks. openpyxl writes no
cached values. openpyxl is installed per user (docs/setup.md); `-I` drops
the user site directory, so it is added back explicitly.
"""

import os
import site
import sys

sys.path.append(site.getusersitepackages())

from openpyxl import Workbook  # noqa: E402
from openpyxl.worksheet.filters import AutoFilter  # noqa: E402
from openpyxl.worksheet.pagebreak import Break  # noqa: E402
from openpyxl.worksheet.properties import Outline  # noqa: E402
from openpyxl.worksheet.table import Table, TableColumn, TableFormula, TableStyleInfo  # noqa: E402

ROWS = [
    ("East", 3, 2.5),
    ("West", 5, 1.5),
    ("East", 1, 10),
    ("North", 2, 4),
    ("South", 6, 0.5),
]


def sales(ws):
    ws.append(["Region", "Qty", "Price", "Amount", "Note"])
    for i, (region, qty, price) in enumerate(ROWS, start=2):
        ws.append([region, qty, price, "=Sales[[#This Row],[Qty]]*Sales[[#This Row],[Price]]", None])
    ws["A7"] = "Total"
    ws["B7"] = "=SUBTOTAL(109,Sales[Qty])"
    ws["D7"] = "=SUBTOTAL(101,Sales[Amount])"
    t = Table(displayName="Sales", ref="A1:E7", totalsRowCount=1)
    t.tableColumns = [TableColumn(id=i + 1, name=n) for i, n in enumerate(["Region", "Qty", "Price", "Amount", "Note"])]
    cols = t.tableColumns
    cols[0].totalsRowLabel = "Total"
    cols[1].totalsRowFunction = "sum"
    cols[3].totalsRowFunction = "average"
    cols[3].calculatedColumnFormula = TableFormula(
        attr_text="Sales[[#This Row],[Qty]]*Sales[[#This Row],[Price]]"
    )
    t.autoFilter = AutoFilter(ref="A1:E6")
    t.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9", showFirstColumn=False, showLastColumn=True, showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(t)

    ws["H1"], ws["I1"] = "Key", "Value"
    ws["H2"], ws["I2"] = "a", 1
    ws["H3"], ws["I3"] = "b", 2
    small = Table(displayName="Small", ref="H1:I3")
    small.tableColumns = [TableColumn(id=1, name="Key"), TableColumn(id=2, name="Value")]
    small.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False, showColumnStripes=True)
    ws.add_table(small)

    ws["G2"] = "=SUM(Sales[Qty])"
    ws["G3"] = "=Sales[[#Totals],[Amount]]"
    ws["G4"] = "=SUM(Small[Value])+ROWS(Sales)"
    ws["G5"] = "=COUNTA(Sales[#Headers])"


def outline(ws):
    ws.sheet_properties.outlinePr = Outline(summaryBelow=False, summaryRight=True)
    ws["A1"] = "Year"
    for r in range(2, 8):
        ws.cell(row=r, column=1, value=f"Item {r}")
        ws.cell(row=r, column=2, value=r * 10)
    ws.row_dimensions.group(2, 6, outline_level=1, hidden=False)
    ws.row_dimensions.group(3, 4, outline_level=2, hidden=True)
    ws.row_dimensions[2].collapsed = True
    ws.column_dimensions.group("C", "D", outline_level=1, hidden=False)
    ws.row_breaks.append(Break(id=10))
    ws.row_breaks.append(Break(id=20))
    ws.col_breaks.append(Break(id=4))


def main(out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    sales(ws)
    outline(wb.create_sheet("Outline"))
    path = os.path.join(out, "provisional-tables.xlsx")
    wb.save(path)
    print(path)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else ".")
