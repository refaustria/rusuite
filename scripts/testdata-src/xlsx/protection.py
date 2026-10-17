"""Generate the PROVISIONAL sheet / workbook protection workbook.

    python3 -I scripts/testdata-src/xlsx/protection.py OUT_DIR

Writes provisional-protection.xlsx (see testdata/README.md): sheet
`Locked` protected with openpyxl's legacy hash of "test" (sorting and
formatting cells allowed, B2 unlocked, C3 a hidden formula), sheet `Sha`
protected with the SHA-512 form of "test" (salt 00..0F, spin count
100000, computed here with hashlib), a hidden and a very hidden sheet,
the second sheet active and the workbook structure locked with the
legacy hash of "wb". openpyxl is installed per user (docs/setup.md);
`-I` drops the user site directory, so it is added back explicitly.
"""

import base64
import hashlib
import site
import sys

sys.path.append(site.getusersitepackages())

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Protection  # noqa: E402
from openpyxl.workbook.protection import WorkbookProtection  # noqa: E402


def sha512(password: str, salt: bytes, spin: int) -> str:
    h = hashlib.sha512(salt + password.encode("utf-16-le")).digest()
    for i in range(spin):
        h = hashlib.sha512(h + i.to_bytes(4, "little")).digest()
    return base64.b64encode(h).decode()


def main(out: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Locked"
    ws["A1"] = "locked"
    ws["B2"] = 5
    ws["B2"].protection = Protection(locked=False)
    ws["C3"] = "=B2*2"
    ws["C3"].protection = Protection(locked=True, hidden=True)
    ws.protection.sheet = True
    ws.protection.password = "test"
    ws.protection.sort = False
    ws.protection.formatCells = False

    sha = wb.create_sheet("Sha")
    sha["A1"] = "x"
    salt = bytes(range(16))
    p = sha.protection
    p.sheet = True
    p.algorithmName = "SHA-512"
    p.saltValue = base64.b64encode(salt).decode()
    p.hashValue = sha512("test", salt, 100000)
    p.spinCount = 100000

    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = 3
    deep = wb.create_sheet("Deep")
    deep.sheet_state = "veryHidden"
    wb.active = 1
    wb.security = WorkbookProtection(workbookPassword="wb", lockStructure=True)
    wb.save(f"{out}/provisional-protection.xlsx")


if __name__ == "__main__":
    main(sys.argv[1])
