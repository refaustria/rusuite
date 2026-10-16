"""Generate the PROVISIONAL notes / threaded comments / hyperlinks workbook.

    python3 -I scripts/testdata-src/xlsx/comments.py OUT_DIR

Writes provisional-comments.xlsx (see testdata/README.md): legacy notes
with authors (one with a custom box size), hyperlinks (web page, e-mail,
place in the workbook with a ScreenTip) and – added by hand because
openpyxl cannot write them – one threaded comment with a reply and one
resolved thread, with Excel's legacy placeholder comments and the
workbook's person list. openpyxl is installed per user (docs/setup.md);
`-I` drops the user site directory, so it is added back explicitly.
"""

import io
import site
import sys
import zipfile

sys.path.append(site.getusersitepackages())

from openpyxl import Workbook  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402

ANN = "{6B8F2E61-3B6E-4D2C-9A55-1F0C2B9D7A01}"
BOB = "{0D5C9B22-8E47-4F61-B3A2-7C1E9D4F2B02}"
T1 = "{A1F0C3D2-1111-4A2B-8C3D-000000000001}"
T1R = "{A1F0C3D2-1111-4A2B-8C3D-000000000002}"
T2 = "{A1F0C3D2-1111-4A2B-8C3D-000000000003}"

PLACEHOLDER = (
    "[Threaded comment]\n\nYour version of Excel allows you to read this threaded comment; "
    "however, any edits to it will get removed if the file is opened in a newer version of Excel. "
    "Learn more: https://go.microsoft.com/fwlink/?linkid=870924\n\nComment:\n    "
)


def fixed(wb):
    """Stable document properties (openpyxl stamps the current time)."""
    import datetime

    t = datetime.datetime(2026, 1, 1)
    wb.properties.created = t
    wb.properties.modified = t
    wb.properties.creator = "Rusuite"
    wb.properties.lastModifiedBy = "Rusuite"


def build():
    wb = Workbook()
    fixed(wb)
    ws = wb.active
    ws.title = "Review"
    for r in [("Item", "Qty", "Link"), ("Pen", 10, "Rusuite"), ("Pad", 4, "Mail us"), ("Ink", 7, "Totals")]:
        ws.append(r)
    ws["A2"].comment = Comment("Ann Lee:\nCheck the supplier price.", "Ann Lee")
    big = Comment("Bob:\nCounted twice.", "Bob")
    big.width = 200
    big.height = 100
    ws["B3"].comment = big
    ws["C2"].hyperlink = "https://example.com/rusuite"
    ws["C3"].hyperlink = "mailto:team@example.com?subject=Order"
    ws["C4"].hyperlink = "#Totals!A1"
    ws["C4"].hyperlink.location = "Totals!A1"
    ws["C4"].hyperlink.target = None
    ws["C4"].hyperlink.tooltip = "Go to the totals"
    # Threaded comments go on A4 and B4: placeholders as Excel writes them.
    ws["A4"].comment = Comment(PLACEHOLDER + "Is this still current?\nReply:\n    Yes, until May.", "tc=" + T1)
    ws["B4"].comment = Comment(PLACEHOLDER + "Fixed.", "tc=" + T2)
    tot = wb.create_sheet("Totals")
    tot["A1"] = "Sum"
    tot["B1"] = "=SUM(Review!B2:B4)"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


THREADED = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<ThreadedComments xmlns="http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments" xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><threadedComment ref="A4" dT="2026-10-01T09:30:00.00" personId="{ANN}" id="{T1}"><text>Is this still current?</text></threadedComment><threadedComment ref="A4" dT="2026-10-02T14:05:00.00" personId="{BOB}" id="{T1R}" parentId="{T1}"><text>Yes, until May.</text></threadedComment><threadedComment ref="B4" dT="2026-10-03T08:00:00.00" personId="{BOB}" id="{T2}" done="1"><text>Fixed.</text></threadedComment></ThreadedComments>"""

PERSONS = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<personList xmlns="http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments" xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><person displayName="Ann Lee" id="{ANN}" userId="Ann Lee" providerId="None"/><person displayName="Bob" id="{BOB}" userId="Bob" providerId="None"/></personList>"""


def add_threads(data):
    """Adds the threaded comment and person parts openpyxl cannot write."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "[Content_Types].xml":
                body = body.replace(
                    b"</Types>",
                    b'<Override PartName="/xl/threadedComments/threadedComment1.xml" '
                    b'ContentType="application/vnd.ms-excel.threadedcomments+xml"/>'
                    b'<Override PartName="/xl/persons/person.xml" '
                    b'ContentType="application/vnd.ms-excel.person+xml"/></Types>',
                )
            elif info.filename == "xl/worksheets/_rels/sheet1.xml.rels":
                body = body.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdTc1" Type="http://schemas.microsoft.com/office/2017/10/relationships/threadedComment" '
                    b'Target="../threadedComments/threadedComment1.xml"/></Relationships>',
                )
            elif info.filename == "xl/_rels/workbook.xml.rels":
                body = body.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdP1" Type="http://schemas.microsoft.com/office/2017/10/relationships/person" '
                    b'Target="persons/person.xml"/></Relationships>',
                )
            z.writestr(info, body)
        for name, text in [
            ("xl/threadedComments/threadedComment1.xml", THREADED),
            ("xl/persons/person.xml", PERSONS),
        ]:
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, text.encode())
    return out.getvalue()


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else "testdata/xlsx"
    with open(f"{out}/provisional-comments.xlsx", "wb") as f:
        f.write(add_threads(build()))


if __name__ == "__main__":
    main()
