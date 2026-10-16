"""Generate the PROVISIONAL chart reference workbooks with openpyxl.

    python3 -I scripts/testdata-src/xlsx/charts.py OUT_DIR

Writes provisional-chart-combo.xlsx, provisional-chart-pie.xlsx and
provisional-chart-mixed.xlsx (see testdata/README.md). openpyxl is installed
per user (`pip3 install --user openpyxl`, see docs/setup.md); `-I` drops the
user site directory, so it is added back explicitly.
"""

import site
import sys

sys.path.append(site.getusersitepackages())

from openpyxl import Workbook  # noqa: E402
from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, PieChart, Reference, ScatterChart, Series  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402


def fixed(wb):
    """Stable document properties (openpyxl stamps the current time)."""
    import datetime

    t = datetime.datetime(2026, 1, 1)
    wb.properties.created = t
    wb.properties.modified = t
    wb.properties.creator = "Rusuite"
    wb.properties.lastModifiedBy = "Rusuite"


def combo(path):
    wb = Workbook()
    fixed(wb)
    ws = wb.active
    ws.title = "Sales"
    rows = [
        ("Month", "Revenue", "Profit"),
        ("Jan", 120, 30),
        ("Feb", 135, 42),
        ("Mar", 150, 38),
        ("Apr", 110, 25),
        ("May", 170, 55),
        ("Jun", 190, 61),
    ]
    for r in rows:
        ws.append(r)
    cats = Reference(ws, min_col=1, min_row=2, max_row=7)
    bar = BarChart()
    bar.type = "col"
    bar.grouping = "clustered"
    bar.title = "Revenue and profit"
    bar.y_axis.title = "EUR (thousands)"
    bar.x_axis.title = "Month"
    bar.add_data(Reference(ws, min_col=2, min_row=1, max_row=7), titles_from_data=True)
    bar.set_categories(cats)
    line = LineChart()
    line.add_data(Reference(ws, min_col=3, min_row=1, max_row=7), titles_from_data=True)
    line.set_categories(cats)
    bar += line
    bar.legend.position = "b"
    ws.add_chart(bar, "E2")
    wb.save(path)


def pie(path):
    wb = Workbook()
    fixed(wb)
    ws = wb.active
    ws.title = "Share"
    rows = [("Region", "Units"), ("North", 40), ("South", 25), ("East", 20), ("West", 15)]
    for r in rows:
        ws.append(r)
    labels = Reference(ws, min_col=1, min_row=2, max_row=5)
    data = Reference(ws, min_col=2, min_row=1, max_row=5)
    p = PieChart()
    p.title = "Units by region"
    p.add_data(data, titles_from_data=True)
    p.set_categories(labels)
    p.dataLabels = DataLabelList()
    p.dataLabels.showPercent = True
    p.dataLabels.showCatName = True
    ws.add_chart(p, "D2")
    d = DoughnutChart()
    d.title = "Doughnut"
    d.holeSize = 60
    d.add_data(data, titles_from_data=True)
    d.set_categories(labels)
    d.legend.position = "r"
    ws.add_chart(d, "D18")
    wb.save(path)


def mixed(path):
    wb = Workbook()
    fixed(wb)
    ws = wb.active
    ws.title = "Data"
    rows = [
        ("X", "Y1", "Y2", "Q1", "Q2"),
        (1, 2.0, 1.5, 10, 4),
        (2, 3.5, 2.5, 12, 6),
        (3, 5.0, 2.0, 9, 8),
        (4, 4.5, 3.5, 14, 5),
        (5, 6.0, 4.0, 11, 7),
    ]
    for r in rows:
        ws.append(r)
    xs = Reference(ws, min_col=1, min_row=2, max_row=6)
    sc = ScatterChart()
    sc.title = "Scatter"
    sc.style = 13
    sc.scatterStyle = "lineMarker"
    for col in (2, 3):
        s = Series(Reference(ws, min_col=col, min_row=1, max_row=6), xs, title_from_data=True)
        s.smooth = True
        sc.series.append(s)
    sc.x_axis.title = "X"
    sc.y_axis.title = "Y"
    sc.y_axis.scaling.min = 0
    sc.y_axis.scaling.max = 8
    sc.y_axis.majorUnit = 2
    ws.add_chart(sc, "G2")

    bar = BarChart()
    bar.type = "bar"
    bar.grouping = "stacked"
    bar.overlap = 100
    bar.title = "Stacked bar"
    bar.add_data(Reference(ws, min_col=4, min_row=1, max_col=5, max_row=6), titles_from_data=True)
    bar.set_categories(xs)
    ws.add_chart(bar, "G18")

    area = AreaChart()
    area.grouping = "stacked"
    area.title = "Stacked area"
    area.add_data(Reference(ws, min_col=4, min_row=1, max_col=5, max_row=6), titles_from_data=True)
    area.set_categories(xs)
    ws.add_chart(area, "P2")

    # Column + line on a secondary axis.
    col = BarChart()
    col.title = "Secondary axis"
    col.add_data(Reference(ws, min_col=4, min_row=1, max_row=6), titles_from_data=True)
    col.set_categories(xs)
    col.y_axis.title = "Q1"
    line = LineChart()
    line.add_data(Reference(ws, min_col=2, min_row=1, max_row=6), titles_from_data=True)
    line.y_axis.axId = 200
    line.y_axis.title = "Y1"
    line.y_axis.crosses = "max"
    col += line
    ws.add_chart(col, "P18")
    wb.save(path)


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else "testdata/xlsx"
    combo(f"{out}/provisional-chart-combo.xlsx")
    pie(f"{out}/provisional-chart-pie.xlsx")
    mixed(f"{out}/provisional-chart-mixed.xlsx")


if __name__ == "__main__":
    main()
