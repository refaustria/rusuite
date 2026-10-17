"""Generate the PROVISIONAL conditional formatting / validation / sort and
filter workbook of P9-C5.

    python3 -I scripts/testdata-src/xlsx/condfmt_c5.py OUT_DIR

Writes provisional-condfmt-c5.xlsx (see testdata/README.md). openpyxl
writes the cells, styles, rules, validations, the AutoFilter and the sort
state; the Excel 2010 parts openpyxl cannot write – the data bar options
(`x14:dataBar`) and the cross-sheet list validation
(`x14:dataValidation`) – are added to the sheet XML afterwards, in the
form Excel writes them. openpyxl is installed per user (docs/setup.md);
`-I` drops the user site directory, so it is added back explicitly.
"""

import datetime
import os
import site
import sys
import zipfile

sys.path.append(site.getusersitepackages())

from openpyxl import Workbook  # noqa: E402
from openpyxl.formatting.rule import (  # noqa: E402
    DataBarRule,
    IconSet,
    FormatObject,
    Rule,
)
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.styles.differential import DifferentialStyle, NumberFormat  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from openpyxl.worksheet.filters import (  # noqa: E402
    ColorFilter,
    DateGroupItem,
    FilterColumn,
    Filters,
    SortCondition,
    SortState,
)

RED = "FFFFC7CE"
X14 = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
XM = "http://schemas.microsoft.com/office/excel/2006/main"
BAR_ID = "{00000001-0000-4000-8000-000000000003}"


def main(out: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rules"
    ws.append(["Date", "Score", "Change", "Trend", "Share"])
    base = datetime.date(2024, 1, 10)
    for i in range(7):
        ws.append([base + datetime.timedelta(days=9 * i), i * 3 + 1, i * 10 - 30, i * 15, i * 40])
    for r in range(2, 9):
        ws.cell(r, 1).number_format = "yyyy-mm-dd"
    red_fill = DifferentialStyle(fill=PatternFill(bgColor=RED), font=Font(color="FF9C0006"))
    ws.conditional_formatting.add(
        "A2:A8",
        Rule(type="timePeriod", timePeriod="lastMonth", dxf=red_fill,
             formula=["AND(MONTH(A2)=MONTH(EDATE(TODAY(),0-1)),YEAR(A2)=YEAR(EDATE(TODAY(),0-1)))"]),
    )
    ws.conditional_formatting.add(
        "B2:B8",
        Rule(type="aboveAverage", aboveAverage=True, stdDev=1, dxf=red_fill),
    )
    ws.conditional_formatting.add(
        "C2:C8",
        DataBarRule(start_type="min", end_type="max", color="FF638EC6", showValue=True,
                    minLength=0, maxLength=100),
    )
    ws.conditional_formatting.add(
        "D2:D8",
        Rule(type="iconSet", iconSet=IconSet(iconSet="3Arrows", cfvo=[
            FormatObject(type="percent", val=0),
            FormatObject(type="percent", val=33, gte=False),
            FormatObject(type="percent", val=67),
        ], reverse=True)),
    )
    # A rule number format (`dxf/numFmt`).
    pts = DifferentialStyle(font=Font(bold=True), numFmt=NumberFormat(numFmtId=164, formatCode='0.0" pts"'))
    ws.conditional_formatting.add(
        "E2:E8",
        Rule(type="cellIs", operator="greaterThan", formula=["100"], dxf=pts),
    )

    lists = wb.create_sheet("Lists")
    for v in ["Apple", "Pear", "Plum"]:
        lists.append([v])
    whole = DataValidation(type="whole", operator="between", formula1="1", formula2="9",
                           errorStyle="warning", showErrorMessage=True, error="1 to 9")
    whole.add("G2:G8")
    custom = DataValidation(type="custom", formula1="H2>G2", showErrorMessage=True)
    custom.add("H2:H8")
    ws.add_data_validation(whole)
    ws.add_data_validation(custom)
    ws["G2"] = 12  # invalid: Circle Invalid Data circles it

    sort = wb.create_sheet("Sort")
    sort.append(["Day", "Amount", "Name"])
    rows = [
        (datetime.date(2024, 1, 5), 10, "Ann"),
        (datetime.date(2024, 1, 20), 20, "Ben"),
        (datetime.date(2024, 2, 3), 30, "Cy"),
        (datetime.date(2025, 3, 1), 40, "Di"),
        (datetime.date(2024, 1, 9), 50, "Ed"),
    ]
    for row in rows:
        sort.append(list(row))
    for r in range(2, 7):
        sort.cell(r, 1).number_format = "yyyy-mm-dd"
    yellow = PatternFill(fill_type="solid", fgColor="FFFFFF00")
    for r in (2, 3, 6):
        sort.cell(r, 2).fill = yellow
    wb._differential_styles.append(DifferentialStyle(fill=PatternFill(bgColor="FFFFFF00")))
    dxf = len(wb._differential_styles.dxf) - 1
    sort.auto_filter.ref = "A1:C6"
    sort.auto_filter.filterColumn.append(FilterColumn(colId=0, filters=Filters(dateGroupItem=[
        DateGroupItem(year=2024, month=1, dateTimeGrouping="month")])))
    sort.auto_filter.filterColumn.append(FilterColumn(colId=1, colorFilter=ColorFilter(dxfId=dxf)))
    sort.auto_filter.sortState = SortState(ref="A2:C6", sortCondition=[
        SortCondition(ref="B2:B6", sortBy="cellColor", dxfId=dxf)])
    # Filtered: row 4 (Feb), row 5 (2025); row 6 (Ed) hidden by hand.
    for r in (4, 5, 6):
        sort.row_dimensions[r].hidden = True

    path = os.path.join(out, "provisional-condfmt-c5.xlsx")
    wb.save(path)
    add_x14(path)


def add_x14(path: str) -> None:
    """Adds the Excel 2010 data bar options and the cross-sheet list."""
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    name = "xl/worksheets/sheet1.xml"
    xml = parts[name].decode()
    bar_ext = (f'<extLst><ext uri="{{B025F937-C7B1-47D3-B67F-A62EFF666E3E}}" xmlns:x14="{X14}">'
               f'<x14:id>{BAR_ID}</x14:id></ext></extLst></cfRule>')
    i = xml.index("<dataBar")
    j = xml.index("</cfRule>", i)
    xml = xml[:j] + bar_ext + xml[j + len("</cfRule>"):]
    ext = (
        f'<extLst><ext uri="{{78C0D931-6437-407d-A8EE-F0AAD7539E65}}" xmlns:x14="{X14}">'
        f'<x14:conditionalFormattings><x14:conditionalFormatting xmlns:xm="{XM}">'
        f'<x14:cfRule type="dataBar" id="{BAR_ID}"><x14:dataBar minLength="0" maxLength="100" gradient="0" '
        'border="1" negativeBarBorderColorSameAsPositive="0" axisPosition="middle">'
        '<x14:cfvo type="autoMin"/><x14:cfvo type="autoMax"/><x14:borderColor rgb="FF638EC6"/>'
        '<x14:negativeFillColor rgb="FF0070C0"/><x14:negativeBorderColor rgb="FF002060"/>'
        '<x14:axisColor rgb="FF7F7F7F"/></x14:dataBar></x14:cfRule><xm:sqref>C2:C8</xm:sqref>'
        '</x14:conditionalFormatting></x14:conditionalFormattings></ext>'
        f'<ext uri="{{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}}" xmlns:x14="{X14}">'
        f'<x14:dataValidations count="1" xmlns:xm="{XM}"><x14:dataValidation type="list" '
        'allowBlank="1" showInputMessage="1" showErrorMessage="1" promptTitle="Fruit" prompt="Pick a fruit">'
        '<x14:formula1><xm:f>Lists!$A$1:$A$3</xm:f></x14:formula1><xm:sqref>F2:F8</xm:sqref>'
        '</x14:dataValidation></x14:dataValidations></ext></extLst></worksheet>'
    )
    xml = xml.replace("</worksheet>", ext)
    parts[name] = xml.encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in parts.items():
            z.writestr(zipfile.ZipInfo(n, date_time=(2026, 10, 17, 0, 0, 0)), data, zipfile.ZIP_DEFLATED)


if __name__ == "__main__":
    main(sys.argv[1])
